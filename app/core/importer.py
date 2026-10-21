"""BOM 文件读取、列映射、清洗、多层展开。"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from .textutil import (
    clean_text,
    is_numeric_text,
    norm_code,
    norm_text,
    parse_number,
)

FIELDS = [
    ("bom_parent", "下级BOM", False),
    ("code", "下级编码", True),
    ("name", "名称", True),
    ("spec", "规格", False),
    ("unit", "单位", False),
    ("qty", "标准用量", True),
    ("loss_rate", "损耗率", False),
    ("alt_text", "替换物料", False),
    ("supply", "获取途径", False),
    ("issue", "发料中心", False),
]
FIELD_KEYS = [f[0] for f in FIELDS]
FIELD_LABELS = {f[0]: f[1] for f in FIELDS}
REQUIRED_FIELDS = {"code", "qty"}

FIELD_KEYWORDS = {
    "bom_parent": ["下级bom", "下级bom编号", "bom层级", "层级", "父项bom", "上级bom", "bom编号", "bom"],
    "code": ["下级编码", "物料编码", "子件编码", "物料号", "料号", "编码", "编号", "code", "partno", "part"],
    "name": ["物料名称", "名称", "品名", "描述", "name", "description"],
    "spec": ["规格型号", "规格/型号", "规格", "型号", "spec", "model"],
    "unit": ["计量单位", "单位", "unit", "uom"],
    "qty": ["标准用量", "单位用量", "基本用量", "用量", "数量", "qty", "quantity", "qtyper"],
    "loss_rate": ["损耗率", "损耗", "loss"],
    "alt_text": ["替换物料", "替代物料", "替换料", "替代料", "替代", "替换", "alt"],
    "supply": ["获取途径", "获取方式", "采购方式", "供应方式", "途径", "source"],
    "issue": ["发料中心", "发料仓", "发料", "领料", "issuing", "issue"],
}

TOTAL_WORDS = ("合计", "小计", "总计", "汇总", "total", "subtotal")
COMMENT_PREFIX = ("#", "//", "备注", "说明", "注:", "注：")


@dataclass
class SheetData:
    path: str
    headers: list[str]
    rows: list[list]
    sheet_names: list[str]
    active_sheet: str
    header_row: int = 0


@dataclass
class CleanOptions:
    strip: bool = True
    upper: bool = True
    skip_empty: bool = True
    skip_total: bool = True
    skip_comment: bool = True
    flatten: bool = False


@dataclass
class CleanedItem:
    row_no: int
    bom_parent: str = ""
    code: str = ""
    name: str = ""
    spec: str = ""
    unit: str = ""
    qty: float | None = None
    loss_rate: str = ""
    alt_text: str = ""
    supply: str = ""
    issue: str = ""
    level: int = 1
    total_qty: float | None = None
    issues: list = field(default_factory=list)


@dataclass
class CleanResult:
    items: list
    skipped: list
    warnings: list[str] = field(default_factory=list)


def _dedupe_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for idx, header in enumerate(headers):
        name = clean_text(header) or f"列{idx + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}({seen[name]})"
        else:
            seen[name] = 0
        result.append(name)
    return result


def _pick_header_row(raw_rows: list[list]) -> int:
    """在前 15 行中找到最像表头的一行。"""
    best_index, best_score = 0, -1
    for idx, row in enumerate(raw_rows[:15]):
        cells = [norm_text(c) for c in row if clean_text(c)]
        if len(cells) < 2:
            continue
        score = len(cells)
        for keyword in sum(FIELD_KEYWORDS.values(), []):
            if any(keyword.replace(" ", "") in cell for cell in cells):
                score += 5
        if score > best_score:
            best_index, best_score = idx, score
    return best_index


def read_file(path: str, sheet: str | None = None) -> SheetData:
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_excel(path, sheet)
    if ext in (".csv", ".txt"):
        return _read_csv(path)
    raise ValueError("仅支持 Excel(.xlsx/.xlsm) 与 CSV 文件")


def _read_excel(path: str, sheet: str | None) -> SheetData:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    ws = wb[sheet] if sheet and sheet in names else wb[names[0]]
    raw = []
    for row in ws.iter_rows(values_only=True):
        raw.append(list(row))
    wb.close()

    header_idx = _pick_header_row(raw)
    if header_idx >= len(raw):
        raise ValueError("文件中没有可用的表头行")
    headers = _dedupe_headers([clean_text(c) for c in raw[header_idx]])
    rows = []
    for row in raw[header_idx + 1 :]:
        row = list(row[: len(headers)])
        row += [None] * (len(headers) - len(row))
        rows.append(row)
    while rows and all(not clean_text(c) for c in rows[-1]):
        rows.pop()
    return SheetData(path, headers, rows, names, ws.title, header_idx)


def _read_csv(path: str) -> SheetData:
    raw_bytes = Path(path).read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "utf-16"):
        try:
            text = raw_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("无法识别 CSV 文件编码（建议另存为 UTF-8）")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = list(csv.reader(text.splitlines(), dialect))
    if not reader:
        raise ValueError("CSV 文件为空")

    header_idx = _pick_header_row(reader)
    headers = _dedupe_headers([clean_text(c) for c in reader[header_idx]])
    rows = []
    for row in reader[header_idx + 1 :]:
        row = list(row[: len(headers)])
        row += [""] * (len(headers) - len(row))
        rows.append(row)
    while rows and all(not clean_text(c) for c in rows[-1]):
        rows.pop()
    name = Path(path).name
    return SheetData(path, headers, rows, [name], name, header_idx)


def auto_mapping(headers: list[str], template: dict | None = None) -> dict:
    """自动猜测字段 → 列名映射；给出模板时优先使用模板。"""
    mapping: dict[str, str] = {}
    used: set[str] = set()

    if template:
        for field_key, header in template.items():
            if field_key in FIELD_KEYS and header in headers and header not in used:
                mapping[field_key] = header
                used.add(header)

    ordered_fields = ["code", "name", "spec", "qty", "unit", "bom_parent",
                      "alt_text", "loss_rate", "supply", "issue"]
    for field_key in ordered_fields:
        if field_key in mapping:
            continue
        best_header, best_score = "", 0
        for header in headers:
            if header in used:
                continue
            hn = norm_text(header)
            score = 0
            for keyword in FIELD_KEYWORDS.get(field_key, []):
                kw = norm_text(keyword)
                if not kw:
                    continue
                if hn == kw:
                    score = max(score, 1000 + len(kw))
                elif kw in hn:
                    score = max(score, 10 + len(kw))
            if score > best_score:
                best_header, best_score = header, score
        if best_header and best_score >= 10:
            mapping[field_key] = best_header
            used.add(best_header)
    return mapping


def mapping_to_index(mapping: dict, headers: list[str]) -> dict[str, int]:
    index = {}
    for field_key, header in mapping.items():
        if header in headers:
            index[field_key] = headers.index(header)
    return index


def _looks_like_total(row: list) -> bool:
    for cell in row[:4]:
        text = norm_text(cell).lower()
        if not text:
            continue
        if any(word in text for word in TOTAL_WORDS):
            return True
    return False


def _looks_like_comment(row: list) -> bool:
    for cell in row[:3]:
        text = clean_text(cell)
        if not text:
            continue
        lower = text.lower()
        if lower.startswith(COMMENT_PREFIX):
            return True
    return False


def clean_rows(
    sheet: SheetData,
    mapping: dict,
    options: CleanOptions | None = None,
) -> CleanResult:
    options = options or CleanOptions()
    index = mapping_to_index(mapping, sheet.headers)
    if "code" not in index:
        raise ValueError("必须映射『下级编码』列")
    if "qty" not in index:
        raise ValueError("必须映射『标准用量』列")

    items: list[CleanedItem] = []
    skipped: list[dict] = []
    warnings: list[str] = []
    if "name" not in index:
        warnings.append("未映射『名称』列，名称+规格模糊匹配将不可用。")

    def cell(row, key):
        if key not in index:
            return ""
        return row[index[key]]

    for offset, row in enumerate(sheet.rows):
        row_no = sheet.header_row + offset + 2
        if options.skip_empty and all(not clean_text(c) for c in row):
            continue
        if options.skip_total and _looks_like_total(row):
            skipped.append({"row_no": row_no, "reason": "合计/小计行", "preview": _preview(row)})
            continue
        if options.skip_comment and _looks_like_comment(row):
            skipped.append({"row_no": row_no, "reason": "注释行", "preview": _preview(row)})
            continue

        def text_of(key):
            return clean_text(cell(row, key))

        code_raw = cell(row, "code")
        code = norm_code(code_raw) if options.upper else clean_text(code_raw).replace(" ", "")
        name = text_of("name")
        spec = text_of("spec")
        unit = text_of("unit")
        bom_parent = clean_text(cell(row, "bom_parent"))
        loss_rate = clean_text(cell(row, "loss_rate"))
        alt_text = clean_text(cell(row, "alt_text"))
        supply = text_of("supply")
        issue = text_of("issue")

        qty_raw = cell(row, "qty")
        qty = parse_number(qty_raw)

        if not code and not name:
            if all(not clean_text(c) for c in row):
                continue
            skipped.append({"row_no": row_no, "reason": "编码与名称均为空", "preview": _preview(row)})
            continue

        issues: list[str] = []
        if not code:
            issues.append("编码为空")
        if not spec:
            issues.append("规格为空")
        if qty is None:
            issues.append("用量无法解析" if clean_text(qty_raw) else "用量为空")
        elif not is_numeric_text(qty_raw):
            warnings.append(f"第{row_no}行用量含非数字字符，已按 {qty} 处理。")

        level = _level_from_parent(bom_parent)
        items.append(
            CleanedItem(
                row_no=row_no,
                bom_parent=norm_text(bom_parent),
                code=code,
                name=name,
                spec=spec,
                unit=unit,
                qty=qty,
                loss_rate=loss_rate,
                alt_text=alt_text,
                supply=supply,
                issue=issue,
                level=level,
                total_qty=qty,
                issues=issues,
            )
        )

    result = CleanResult(items=items, skipped=skipped, warnings=warnings)
    if options.flatten and items:
        flatten_items(items)
    return result


def _preview(row: list) -> str:
    parts = [clean_text(c) for c in row if clean_text(c)]
    return " | ".join(parts[:5])[:120]


def _level_from_parent(text: str) -> int:
    text = clean_text(text)
    if not text:
        return 1
    if re.fullmatch(r"\d+(?:\.\d+)*", text):
        return text.count(".") + 1
    return 1


def flatten_items(items: list[CleanedItem]) -> list[str]:
    """按 下级BOM 层级路径递归展开，总用量=各级标准用量逐层相乘，含循环检测。"""
    by_parent: dict[str, CleanedItem] = {}
    by_code: dict[str, CleanedItem] = {}
    for item in items:
        key = norm_text(item.bom_parent)
        if key and key not in by_parent:
            by_parent[key] = item
        code = norm_code(item.code)
        if code and code not in by_code:
            by_code[code] = item

    def find_parent(item: CleanedItem):
        key = norm_text(item.bom_parent)
        if not key:
            return None
        if "." in key:
            prefix = key.rsplit(".", 1)[0]
            if prefix in by_parent:
                return by_parent[prefix]
        if key in by_parent and by_parent[key] is not item:
            return by_parent[key]
        if key in by_code and by_code[key] is not item:
            return by_code[key]
        return None

    totals: dict[int, float] = {}
    levels: dict[int, int] = {}
    cycle_items: list[str] = []

    def resolve(item: CleanedItem, stack: set[int]) -> float:
        key = id(item)
        if key in totals:
            return totals[key]
        if key in stack:
            raise RecursionError("循环嵌套")
        stack.add(key)
        qty = float(item.qty or 0.0)
        parent = find_parent(item)
        if parent is None:
            total = qty
            level = 1
        else:
            total = qty * resolve(parent, stack)
            level = levels.get(id(parent), 1) + 1
        stack.discard(key)
        totals[key] = total
        levels[key] = level
        return total

    for item in items:
        try:
            item.total_qty = round(resolve(item, set()), 6)
            item.level = levels.get(id(item), item.level)
        except RecursionError:
            item.total_qty = item.qty
            item.level = 1
            if "BOM 循环嵌套" not in item.issues:
                item.issues.append("BOM 循环嵌套")
            cycle_items.append(item.code or f"第{item.row_no}行")
    return cycle_items


def collect_alt_pairs(items: list[CleanedItem]) -> list[tuple[list[str], str, str]]:
    """从明细中抽取替换物料关系，供替代组自动生成。"""
    from .textutil import parse_alt_codes

    pairs = []
    for item in items:
        alts = [norm_code(c) for c in parse_alt_codes(item.alt_text) if c]
        alts = [a for a in alts if a]
        if not alts:
            continue
        codes = ([item.code] if item.code else []) + alts
        if len(codes) >= 2:
            pairs.append((codes, item.name, item.spec))
    return pairs
