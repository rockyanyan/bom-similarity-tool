"""Excel 差异报告导出。"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import row_to_text
from .textutil import format_qty

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="1F3864")
LABEL_FONT = Font(name="微软雅黑", size=10, bold=True, color="404040")
BODY_FONT = Font(name="微软雅黑", size=10)
ALERT_FONT = Font(name="微软雅黑", size=10, bold=True, color="9C0006")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TYPE_FILLS = {
    "完全匹配": PatternFill("solid", fgColor="F2F2F2"),
    "等价替换匹配": PatternFill("solid", fgColor="DDEBF7"),
    "数量变更": PatternFill("solid", fgColor="FFF2CC"),
    "新增物料": PatternFill("solid", fgColor="E2EFDA"),
    "缺失物料": PatternFill("solid", fgColor="FCE4D6"),
}
ALERT_FILL = PatternFill("solid", fgColor="FFC7CE")

DETAIL_COLUMNS = [
    ("差异类型", 14),
    ("层级", 6),
    ("下级BOM", 10),
    ("库内编码", 16),
    ("新BOM编码", 16),
    ("名称", 22),
    ("规格", 24),
    ("单位", 6),
    ("库内标准用量", 12),
    ("新BOM标准用量", 13),
    ("用量变化", 10),
    ("变化率", 9),
    ("权重", 7),
    ("库内获取途径", 12),
    ("新BOM获取途径", 12),
    ("库内发料中心", 12),
    ("新BOM发料中心", 12),
    ("替换物料", 16),
    ("附加差异", 16),
    ("说明", 22),
]


def export_report(result, path: str, options=None) -> str:
    wb = Workbook()
    _build_summary(wb.active, result, options)
    _build_detail(wb.create_sheet("差异明细"), result)
    wb.save(path)
    return path


def _build_summary(ws, result, options=None):
    ws.title = "汇总"
    ws.merge_cells("A1:D1")
    ws["A1"] = "BOM 相似度对比报告"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30

    rows = [
        ("对比方案 A（库内）", result.label_a),
        ("对比方案 B（新 BOM）", result.label_b),
        ("对比时间", result.time),
        ("相似度算法", result.algorithm_label),
        ("相似度得分", f"{result.score:.2f}%"),
        ("物料数 A / B", f"{result.total_a} / {result.total_b}"),
        ("匹配成功物料数", result.matched),
        ("高权重差异项", result.high_weight_count),
    ]
    if options is not None:
        rows += [
            ("模糊匹配阈值", f"{options.fuzzy_threshold:.2f}"),
            ("高权重阈值", f"{options.high_weight:.2f}"),
            ("获取途径差异", "已标记" if options.diff_supply else "未标记"),
            ("发料中心差异", "已标记" if options.diff_issue else "未标记"),
        ]

    start = 3
    for offset, (label, value) in enumerate(rows):
        row = start + offset
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = BODY_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        if label == "相似度得分":
            cell.font = Font(name="微软雅黑", size=14, bold=True, color="1F3864")

    stat_start = start + len(rows) + 1
    ws.cell(row=stat_start, column=1, value="差异分类统计").font = TITLE_FONT
    ws.cell(row=stat_start, column=1).font = Font(name="微软雅黑", size=12, bold=True, color="1F3864")
    header_row = stat_start + 1
    for col, name in enumerate(("差异类型", "数量", "占比"), start=1):
        cell = ws.cell(row=header_row, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

    total = max(1, len(result.rows))
    for offset, (dtype, count) in enumerate(result.stats.items()):
        row = header_row + 1 + offset
        ws.cell(row=row, column=1, value=dtype).font = BODY_FONT
        ws.cell(row=row, column=2, value=count).font = BODY_FONT
        ws.cell(row=row, column=3, value=f"{count / total * 100:.1f}%").font = BODY_FONT
        for col in range(1, 4):
            ws.cell(row=row, column=col).border = BORDER
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="center")
        if dtype in TYPE_FILLS:
            ws.cell(row=row, column=1).fill = TYPE_FILLS[dtype]

    note_row = header_row + len(result.stats) + 2
    ws.cell(row=note_row, column=1, value="复核提示").font = LABEL_FONT
    notes = list(result.warnings) or ["无"]
    notes.append("相似度分数仅供参考，等价物料判定结果需人工复核。")
    for offset, text in enumerate(notes):
        ws.cell(row=note_row + 1 + offset, column=1, value=f"· {text}").font = BODY_FONT

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 12
    ws.sheet_view.showGridLines = False


def _build_detail(ws, result):
    headers = [name for name, _ in DETAIL_COLUMNS]
    for idx, (name, width) in enumerate(DETAIL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 26

    for row_idx, row in enumerate(result.rows, start=2):
        text = row_to_text(row)
        for col_idx, header in enumerate(headers, start=1):
            value = text.get(header, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = ALERT_FONT if row["highlight"] else BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if col_idx in (2, 3, 8, 9, 10, 11, 12, 13) else "left",
                                       vertical="center", wrap_text=False)
            cell.fill = ALERT_FILL if row["highlight"] else TYPE_FILLS.get(row["type"], TYPE_FILLS["完全匹配"])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(result.rows) + 1)}"


def export_scheme_items(scheme: dict, items: list, path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM明细"
    headers = [
        ("行号", 6), ("层级", 6), ("下级BOM", 10), ("下级编码", 18), ("名称", 24),
        ("规格", 26), ("单位", 6), ("标准用量", 11), ("展开总用量", 12),
        ("损耗率", 8), ("替换物料", 18), ("获取途径", 12), ("发料中心", 12), ("异常标记", 20),
    ]
    for idx, (name, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row_idx, item in enumerate(items, start=2):
        values = [
            item.get("row_no"), item.get("level"), item.get("bom_parent"), item.get("code"),
            item.get("name"), item.get("spec"), item.get("unit"), item.get("qty"),
            item.get("total_qty"), item.get("loss_rate"), item.get("alt_text"),
            item.get("supply"), item.get("issue"), "；".join(item.get("issues") or []),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 3, 7, 8, 9, 10) else "left",
                                       vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:N{max(1, len(items) + 1)}"
    wb.save(path)
    return path
